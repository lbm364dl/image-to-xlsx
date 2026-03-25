import pymupdf
import pypdfium2
from openpyxl import Workbook
from utils import file_extension
from PIL import Image
from io import BytesIO


class Document:
    def __init__(
        self,
        document,
        fixed_decimal_places=0,
        method="surya",
    ):
        self.method = method
        self.fixed_decimal_places = fixed_decimal_places
        self.extension = file_extension(document["name"])
        self.tot_pages = 1
        self.pdf_bytes = None

        if self.extension == "pdf":
            self.pdf = pymupdf.open("pdf", document["content"])
            self.tot_pages = self.pdf.page_count
            self.pdf_bytes = document["content"]

        self.set_page_nums(document["pages"])
        self.load_pages(document["content"], method)

        self.workbook = Workbook()
        self.workbook.remove(self.workbook.active)
        self.footers_workbook = Workbook()
        self.footers_workbook.active.append([
            "page_number",
            "table_number",
            "footer_text",
        ])

    def set_page_nums(self, page_ranges):
        self.page_nums = set()
        for start, end in page_ranges:
            self.page_nums |= set(range(start, min(self.tot_pages, end) + 1))

    def load_pages(self, document, method):
        if self.extension == "pdf":
            if method == "pdf-text":
                # pdf-text extraction is handled by the service using raw PDF bytes;
                # pages dict maps page numbers to None (no images needed).
                self.pages = {i: None for i in self.page_nums}
            else:
                pdf = pypdfium2.PdfDocument(BytesIO(document))
                if hasattr(pdf, "render"):
                    self.pages = dict(
                        zip(
                            self.page_nums,
                            pdf.render(
                                pypdfium2.PdfBitmap.to_pil,
                                page_indices=[i - 1 for i in self.page_nums],
                                scale=2,
                            ),
                        )
                    )
                else:
                    pages = {}
                    for page_num in self.page_nums:
                        page = pdf.get_page(page_num - 1)
                        try:
                            bitmap = page.render(scale=2)
                            pages[page_num] = bitmap.to_pil()
                        finally:
                            page.close()
                    self.pages = pages
        else:
            from PIL import ImageOps
            self.pages = {1: ImageOps.exif_transpose(Image.open(BytesIO(document)))}
