"""Shared fixtures for the test suite."""

import sys
from pathlib import Path

import pytest

# Make sure the main app source is importable
SRC = Path(__file__).resolve().parent.parent / "src" / "image-to-xlsx"
if str(SRC) not in sys.path:
    sys.path.insert(0, str(SRC))

INPUTS_DIR = SRC / "inputs"


@pytest.fixture
def sample_service_cells():
    """A realistic extraction-service response (cells list)."""
    return [
        {"row": 0, "col": 0, "text": "Year", "confidence": 99.0},
        {"row": 0, "col": 1, "text": "Population", "confidence": 97.5},
        {"row": 0, "col": 2, "text": "GDP", "confidence": 98.0},
        {"row": 1, "col": 0, "text": "2020", "confidence": 95.0},
        {"row": 1, "col": 1, "text": "331,002,651", "confidence": 88.0},
        {"row": 1, "col": 2, "text": "20,936.60", "confidence": 91.0},
        {"row": 2, "col": 0, "text": "2021", "confidence": 96.0},
        {"row": 2, "col": 1, "text": "332,403,650", "confidence": 85.0},
        {"row": 2, "col": 2, "text": "23,315.08", "confidence": 90.0},
    ]


@pytest.fixture
def sample_service_response(sample_service_cells):
    """A full extraction-service JSON response."""
    return {
        "tables": [
            {
                "bbox": [10, 20, 500, 400],
                "cells": sample_service_cells,
            }
        ]
    }


@pytest.fixture
def empty_service_response():
    """An extraction-service response with no tables."""
    return {"tables": []}


@pytest.fixture
def multi_table_response():
    """Service response with two tables."""
    return {
        "tables": [
            {
                "bbox": [0, 0, 100, 50],
                "cells": [
                    {"row": 0, "col": 0, "text": "A", "confidence": 99},
                    {"row": 0, "col": 1, "text": "B", "confidence": 98},
                ],
            },
            {
                "bbox": [0, 60, 100, 120],
                "cells": [
                    {"row": 0, "col": 0, "text": "X", "confidence": None},
                    {"row": 1, "col": 0, "text": "Y", "confidence": None},
                ],
            },
        ]
    }


@pytest.fixture
def make_document():
    """Factory that creates a minimal Document from raw bytes."""

    def _make(content, name="test.png", method="surya"):
        from document import Document

        doc_dict = {"name": name, "content": content, "pages": [(1, 1)]}
        return Document(doc_dict, fixed_decimal_places=0, method=method)

    return _make


@pytest.fixture
def stub_page():
    """Create a Page-like object sufficient for Table operations."""

    class _Doc:
        fixed_decimal_places = 0
        method = "surya"

        def __init__(self):
            from openpyxl import Workbook

            self.workbook = Workbook()
            self.workbook.remove(self.workbook.active)
            self.footers_workbook = Workbook()
            self.footers_workbook.active.append(
                ["page_number", "table_number", "footer_text"]
            )

    class _Page:
        page_num = 1
        document = _Doc()

    return _Page()
